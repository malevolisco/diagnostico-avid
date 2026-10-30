#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Comprueba el resultado de una ejecución del programa sobre el disco de prueba.

Uso:  python tests/comprobar_resultado.py CARPETA_INFORME CARPETA_BRUTOS [--ffprobe RUTA] [--reanudado] [--agrupar proyecto_mes|proyecto]
Sale con código 1 si algo no cuadra, para que el flujo de GitHub falle.
--reanudado: tras una segunda ejecución sobre la misma carpeta de brutos, todos
los proyectos deben venir marcados como reanudados (no se han vuelto a exportar).
"""

import argparse
import glob
import json
import shutil
import subprocess
import sys
from pathlib import Path


sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # la consola de Windows (cp1252) no tiene chino

def fallo(msg):
    print("FALLO:", msg)
    sys.exit(1)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("informe")
    ap.add_argument("brutos")
    ap.add_argument("--ffprobe", default=shutil.which("ffprobe"))
    ap.add_argument("--reanudado", action="store_true")
    ap.add_argument("--rdd9", action="store_true", help="los brutos tienen que estar unidos con bmx en RDD 9 (bmx junto al programa)")
    ap.add_argument("--agrupar", choices=["proyecto_mes", "proyecto"], default="proyecto_mes",
                    help="proyecto_mes = config por defecto (XDCAM HD 50, sin importados) | proyecto = agrupar proyecto, formato auto, con importados")
    a = ap.parse_args()

    carpetas = glob.glob(str(Path(a.informe) / "Diagnostico_*"))
    carpetas = [c for c in carpetas if Path(c).is_dir()]
    if not carpetas:
        fallo("no hay carpeta Diagnostico_* en %s" % a.informe)
    carpeta = Path(sorted(carpetas)[-1])
    for f in ["informe.html", "resumen.json", "bins.json", "mxf.json", "clips.json", "cruce.json", "inventario.csv", "diagnostico.log"]:
        if not (carpeta / f).exists():
            fallo("falta %s" % f)
    r = json.loads((carpeta / "resumen.json").read_text(encoding="utf-8"))
    # clips_fisicos: 9 + 2 NTSC + 3 APROX + agencia + pieza
    # + importado + render sin bin + 4 de China. huerfanos: HUERFANO, SOLOAUDIO, CUARENTENA y los 2 "APROX B" (el render no cuenta).
    # China: 2 bins, uno en una ruta de más de 260 caracteres (en Windows hay que leerlo con \\?\).
    # ENLACES: 2 ficheros enlazados (AMA) fuera de Avid MediaFiles y 1 clip con la media en otra carpeta (--media-extra).
    # ORIGEN: 2 bins (uno "compras"), un MOV enlazado de una Panasonic y 2 clips importados sin pistas de origen.
    # offline: CLIP OFFLINE, los 2 "APROX B" ambiguos y VIEJO OMF (su media es OMF, que no se lee)
    # DUPLICADO: una entrevista DV y un trozo suyo copiado a otro fichero (mismo subcódigo de cámara y mismos fotogramas),
    # y una entrevista HD y un trozo suyo (mismo timecode del fichero y mismo sonido, muestra a muestra)
    esperado = {"proyectos": 8, "bins": 12, "bins_ilegibles": 1, "offline": 4, "renders_en_bins": 1, "clips_compartidos": 1, "clips_fisicos": 32,
                "huerfanos": 5, "trozos_repetidos": 2}
    for k, v in esperado.items():
        if r.get(k) != v:
            fallo("resumen.%s = %r, esperado %r" % (k, r.get(k), v))
    modos = r.get("modo_coincidencia") or {}
    if not modos.get("le"):
        fallo("el cruce bin<->MXF por identificador no ha coincidido con nada: %r" % modos)
    if modos.get("enlace") != 3:
        fallo("media enlazada (AMA): %r clips enlazados encontrados, esperados 3" % modos.get("enlace"))
    if json.loads((carpeta / "cruce.json").read_text(encoding="utf-8")).get("offline_en_omf") != 1 \
            or "media en OMF" not in (carpeta / "informe.html").read_text(encoding="utf-8"):
        fallo("OMF: el clip VIEJO OMF debe salir offline marcado 'media en OMF' en cruce.json y en el informe")
    if modos.get("nombre_duracion") != 1:
        fallo("cruce aproximado: nombre_duracion = %r, esperado 1" % modos.get("nombre_duracion"))
    cruce = json.loads((carpeta / "cruce.json").read_text(encoding="utf-8"))
    aprox = next((p for p in cruce["por_proyecto"] if p["proyecto"] == "APROX"), {})
    if (aprox.get("clips_aproximados"), aprox.get("clips_ambiguos")) != (1, 2):
        fallo("APROX: aproximados %r, ambiguos %r; esperado 1 y 2" % (aprox.get("clips_aproximados"), aprox.get("clips_ambiguos")))
    grab = (json.loads((carpeta / "cruce.json").read_text(encoding="utf-8")).get("fecha_grabacion") or {})
    # NTSC (columnas Shoot Date/Time del bin) y el de ORIGEN con el reloj desfasado; la Panasonic del 1/1/2000, descartada;
    # y la entrevista DV de DUPLICADO y su trozo, por el subcódigo de la cámara (13/11/2002: no cuenta como desfasado)
    if (grab.get("clips"), sum((grab.get("imposibles") or {}).values()), grab.get("desfasados")) != (5, 1, 1):
        fallo("fecha de grabación: %r clips, %r imposibles, %r desfasados; esperados 5, 1 y 1"
              % (grab.get("clips"), grab.get("imposibles"), grab.get("desfasados")))
    if r.get("mxf_error"):
        fallo("%d MXF con error de lectura" % r["mxf_error"])
    no_legibles = [x["ruta"].replace("\\", "/") for x in r.get("no_legibles", [])]
    if no_legibles != ["ABORTO/Bin corrupto.avb"]:
        fallo("no legibles %r, esperado solo el bin corrupto" % no_legibles)
    clips_china = {c["nombre"]: c["origen"] for c in json.loads((carpeta / "clips.json").read_text(encoding="utf-8"))
                   if c["nombre"] in ("entrevista 北京", "DJI_0012", "C0380", "新闻 CCTV 0916")}
    if clips_china != {"entrevista 北京": "Cámara (fichero importado)", "DJI_0012": "Cámara (fichero importado)",
                       "C0380": "Cámara (fichero importado)", "新闻 CCTV 0916": "Agencia (bin de agencias)"}:
        fallo("China: nombres u origen %r (nombres en chino leídos del bin, cámara propia y bin de agencias)" % clips_china)
    if not glob.glob(str(carpeta / "miniaturas" / "*.jpg")):
        fallo("no hay miniaturas")
    print("Análisis OK: %d proyectos, %d clips físicos, cruce %s" % (r["proyectos"], r["clips_fisicos"], r["modo_coincidencia"]))

    res_path = Path(a.brutos) / "brutos_resultado.json"
    if not res_path.exists():
        fallo("no hay brutos_resultado.json")
    res = json.loads(res_path.read_text(encoding="utf-8"))
    por = {b["proyecto"]: b for b in res}
    if a.agrupar == "proyecto_mes":
        # Lo que pidió el Archivo (oct. 2026): por proyecto, PROYECTO_AAAA-MM, todo convertido a XDCAM HD 50 1080i25 16:9.
        # Fuera la agencia (SISMO, por la ruta del bin) y el vídeo de móvil importado. La pieza, aparte. NTSC entra:
        # sus MXF se importaron de una tarjeta (CONTENTS) aunque el origen diga "Fichero importado".
        # SIN_PROYECTO: huérfano 2 s + cuarentena 1 s + los 2 "APROX B" ambiguos de 2 s. Renders de efectos, fuera.
        # El bruto de piezas (15 SEPT CENTROAMERICA PIEZAS SEPTIEMBRE DE 2021) no se exporta: exportar_piezas = false
        # China: DNxHD 1080i50 recodificado sin separar campos; el dron y la A7 importados de una carpeta del PC son material;
        # el vídeo del bin de agencias, fuera.
        esperado_brutos = {"15 SEPT CENTROAMERICA SEPTIEMBRE DE 2021": ("convertir", 3, 7.0),
                           "20260916_CHINA 北京 SEPTIEMBRE DE 2021": ("convertir", 3, 6.0),
                           # ENLACES: vertical de móvil con estéreo y HDR HLG enlazados (AMA) y el clip del otro disco
                           "ENLACES SEPTIEMBRE DE 2021": ("convertir", 3, 6.0),
                           # ORIGEN: la Panasonic (propio por los metadatos) y los dos dudosos, que entran marcados
                           "ORIGEN SEPTIEMBRE DE 2021": ("convertir", 3, 6.0),
                           "ABORTO SEPTIEMBRE DE 2021": ("convertir", 2, 4.0), "APROX SEPTIEMBRE DE 2021": ("convertir", 1, 2.0),
                           "NTSC SEPTIEMBRE DE 2021": ("convertir", 2, 4.0), "SIN PROYECTO SEPTIEMBRE DE 2021": ("convertir", 4, 7.0),
                           # DUPLICADO: solo las dos entrevistas; los trozos repetidos, fuera (en la papelera del montaje). La DV
                           # es de 2002 (fecha de la cámara) y la HD de 2021: más de separar_si_pasan_dias, en brutos aparte
                           "DUPLICADO NOVIEMBRE DE 2002": ("convertir", 1, 6.0), "DUPLICADO SEPTIEMBRE DE 2021": ("convertir", 1, 6.0)}
    else:
        # Un bruto por tarjeta (_Tnn_fecha), y aparte agencias, importados y piezas. 15 SEPT: el clip "2017 07 13" es del día 10
        # (otra tarjeta); el de Reuters lo dice la ruta del bin; la pieza se llama como la secuencia. El vídeo de móvil
        # importado va a _IMPORTADOS. Los renders de efectos no entran. NTSC: mismo día pero dos carpetas de tarjeta.
        # SIN_PROYECTO: huérfano 2 s + cuarentena 1 s + los 2 "APROX B" ambiguos de 2 s.
        esperado_brutos = {"15 SEPT CENTROAMERICA_T01_2021-09-10": ("copia", 1, 2.0), "15 SEPT CENTROAMERICA_T02_2021-09-15": ("copia", 2, 5.0),
                           "15 SEPT CENTROAMERICA_AGENCIAS": ("copia", 1, 2.0),
                           "15 SEPT CENTROAMERICA_IMPORTADOS": ("copia", 1, 2.0),
                           # China: el dron y la A7 (importados de una carpeta del PC, material propio) son una tanda de volcado;
                           # la entrevista, de la tarjeta J: (XDROOT). DNxHD 1080i50 igual en todo: copia sin recomprimir.
                           "20260916_China 北京_T01_2021-09-15": ("copia", 2, 4.0), "20260916_China 北京_T02_2021-09-15": ("copia", 1, 2.0),
                           "20260916_China 北京_AGENCIAS": ("copia", 1, 2.0),
                           # ENLACES: el del otro disco (DNxHD, se copia) es una tanda sin carpeta de tarjeta; el móvil y la GoPro
                           # comparten carpeta de tarjeta (E:\RODAJE) y se convierten (formatos distintos)
                           "ENLACES_T01_2021-09-15": ("copia", 1, 2.0), "ENLACES_T02_2021-09-15": ("convertir", 2, 4.0),
                           # ORIGEN: los dudosos entran con el material (no van a _IMPORTADOS); la Panasonic es H.264, se convierte
                           "ORIGEN_T01_2021-09-15": ("convertir", 3, 6.0),
                           "ABORTO_T01_2021-09-15": ("convertir", 2, 4.0), "APROX_T01_2021-09-15": ("copia", 1, 2.0),
                           "NTSC_T01_2021-09-15": ("copia", 1, 2.0), "NTSC_T02_2021-09-15": ("copia", 1, 2.0),
                           "SIN_PROYECTO_T01_2021-09-15": ("copia", 4, 7.0),
                           # DUPLICADO: las dos entrevistas (DV NTSC y HD: formatos distintos, se convierte); los trozos repetidos, fuera
                           "DUPLICADO_T01_2021-09-15": ("convertir", 2, 12.0)}
    if sorted(por) != sorted(esperado_brutos):
        fallo("brutos %s, esperados %s" % (sorted(por), sorted(esperado_brutos)))
    for proy, (modo, incl, seg) in esperado_brutos.items():
        b = por.get(proy)
        if not b:
            fallo("falta el bruto de %s" % proy)
        if b["modo"] != modo or b["incluidos"] != incl or abs(b["segundos"] - seg) > 0.2 or not b["estado"].startswith("OK"):
            fallo("bruto %s: %s" % (proy, {k: b[k] for k in ["modo", "incluidos", "segundos", "estado"]}))
        mxf = Path(a.brutos) / b["fichero"]
        if not mxf.exists() or not (Path(a.brutos) / b["shotlist"]).exists() or not (Path(a.brutos) / b["contactos"]).exists():
            fallo("faltan ficheros del bruto %s" % proy)
        if a.ffprobe:
            out = subprocess.run([a.ffprobe, "-v", "error", "-show_entries", "stream=codec_type", "-of", "default=nw=1:nk=1", str(mxf)],
                                 capture_output=True, text=True).stdout.split()
            if out.count("video") != 1 or out.count("audio") != b["pistas_audio"]:
                fallo("bruto %s: pistas %s, esperadas 1 vídeo + %d audio" % (proy, out, b["pistas_audio"]))
            tc = subprocess.run([a.ffprobe, "-v", "error", "-show_entries", "format_tags=timecode", "-of", "default=nw=1:nk=1", str(mxf)],
                                capture_output=True, text=True).stdout.strip()
            sep = ";" if proy.startswith(("NTSC", "DUPLICADO")) and a.agrupar == "proyecto" else ":"  # drop-frame solo si sale a 29,97
            if len(tc) != 11 or tc[8] != sep:
                fallo("bruto %s: timecode %r, esperado con '%s' antes de los frames" % (proy, tc, sep))
            if a.agrupar == "proyecto_mes":
                v = subprocess.run([a.ffprobe, "-v", "error", "-select_streams", "v:0", "-show_entries",
                                    "stream=codec_name,profile,width,height,pix_fmt,field_order,r_frame_rate,display_aspect_ratio",
                                    "-of", "default=nw=1", str(mxf)], capture_output=True, text=True).stdout.split()
                v = dict(x.split("=", 1) for x in v if "=" in x)
                xdcam = {"codec_name": "mpeg2video", "profile": "4:2:2", "width": "1920", "height": "1080", "pix_fmt": "yuv422p",
                         "field_order": "tt", "r_frame_rate": "25/1", "display_aspect_ratio": "16:9"}
                if any(v.get(k) != x for k, x in xdcam.items()):
                    fallo("bruto %s: vídeo %s, esperado XDCAM HD 50 1080i50 16:9 %s" % (proy, v, xdcam))
                acabado_profesional(proy, b, mxf, tc, a)
        tcs, _, clips, fechas = tc_shotlist(Path(a.brutos) / b["shotlist"])
        if a.agrupar == "proyecto_mes":
            if fechas != sorted(fechas):
                fallo("shotlist %s: no está en orden cronológico (grabación o, si no consta, fichero): %s" % (proy, fechas))
            # NTSC: los dos se volcaron en el mismo segundo; las columnas "Shoot Date/Time" del bin dicen que NTSC 2 se grabó antes
            if proy.startswith("NTSC") and clips != ["NTSC 2", "NTSC 1"]:
                fallo("shotlist %s: orden %s, esperado por hora de grabación ['NTSC 2', 'NTSC 1']" % (proy, clips))
            if {"SISMO 15 SEPT", "MOVIL 15 SEPT", "新闻 CCTV 0916"} & set(clips):
                fallo("shotlist %s: lleva importados o agencias (%s)" % (proy, clips))
            if "CHINA" in proy and sorted(clips) != ["C0380", "DJI_0012", "entrevista 北京"]:
                fallo("shotlist %s: clips %s, esperados la entrevista, el dron y la A7" % (proy, clips))
            if proy.startswith("ENLACES"):
                enlaces(proy, Path(a.brutos) / b["shotlist"], mxf, clips, a)
            if proy.startswith("ORIGEN"):
                origen_por_senales(proy, Path(a.brutos) / b["shotlist"])
        sep = ";" if proy.startswith(("NTSC", "DUPLICADO")) and a.agrupar == "proyecto" else ":"
        if not tcs or any(len(t) != 11 or t[8] != sep for t in tcs):
            fallo("shotlist %s: timecodes %r, esperado '%s' antes de los frames" % (proy, tcs, sep))
    if a.reanudado:
        no = [b["proyecto"] for b in res if not b.get("reanudado")]
        if no:
            fallo("segunda ejecución: se han vuelto a exportar %s en vez de reanudar" % no)
    print("Brutos OK%s:" % (" (reanudados)" if a.reanudado else ""), ", ".join("%s (%s, %.1fs)" % (p, por[p]["modo"], por[p]["segundos"]) for p in esperado_brutos))


def acabado_profesional(proy, b, mxf, tc, a):
    """Ficha técnica conforme, MD5 correcto, TC interno del vídeo igual al del fichero y fecha de creación."""
    import hashlib
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    import mxf_util
    if b.get("conforme") is not True or not (Path(a.brutos) / str(b.get("ficha"))).exists():
        fallo("bruto %s: ficha técnica %r, conforme %r" % (proy, b.get("ficha"), b.get("conforme")))
    if a.rdd9:
        ficha = (Path(a.brutos) / str(b.get("ficha"))).read_text(encoding="utf-8")
        if b.get("rdd9") is not True or "OK     Envoltura SMPTE RDD 9 (BBC bmx" not in ficha:
            fallo("bruto %s: no está en RDD 9 (rdd9 %r; ficha: %s)" % (proy, b.get("rdd9"),
                                                                     [l for l in ficha.splitlines() if "RDD 9" in l]))
    md5 = mxf.with_suffix(".md5")
    if not md5.exists() or md5.read_text(encoding="utf-8") != "%s *%s\n" % (hashlib.md5(mxf.read_bytes()).hexdigest(), mxf.name):
        fallo("bruto %s: el .md5 falta o no cuadra" % proy)
    gop = mxf_util.leer_tc_gop(mxf, 1)
    if gop != [tc]:
        fallo("bruto %s: TC interno %s, TC del fichero %s" % (proy, gop, tc))
    creado = mxf_util.fecha_creacion(mxf) or ""
    if not creado.startswith("20"):
        fallo("bruto %s: sin fecha de creación (%r)" % (proy, creado))


def origen_por_senales(proy, shotlist):
    """ORIGEN: la Panasonic con un nombre que no está en ninguna lista entra como propia por los metadatos; lo importado
    sin pistas y lo del bin "compras" entran marcados "para revisar", con el motivo."""
    from openpyxl import load_workbook
    filas = list(load_workbook(shotlist, read_only=True).active.iter_rows(values_only=True))
    cab = filas[0]
    fila = {f[cab.index("Clip")]: dict(zip(cab, f)) for f in filas[1:]}
    esperado = {"ENTREVISTA ALCALDE": "Material propio", "RECURSO SIN PISTAS": "Dudoso · para revisar",
                "VISTA AEREA": "Dudoso · para revisar"}
    visto = {k: (fila.get(k) or {}).get("Origen") for k in esperado}
    if visto != esperado:
        fallo("shotlist %s: origen %r, esperado %r" % (proy, visto, esperado))
    if "Panasonic" not in str(fila["ENTREVISTA ALCALDE"].get("Por qué")) or "compras" not in str(fila["VISTA AEREA"].get("Por qué")):
        fallo("shotlist %s: motivos %r" % (proy, {k: v.get("Por qué") for k, v in fila.items()}))
    # Fechas: la del 1/1/2000 no se usa (fecha de grabación vacía, aviso); la de 26 h tras el volcado sí, con aviso
    alcalde, recurso = fila["ENTREVISTA ALCALDE"], fila["RECURSO SIN PISTAS"]
    if alcalde.get("Fecha grabación") or "imposible" not in str(alcalde.get("Aviso fecha")) \
            or not recurso.get("Fecha grabación") or "26 h" not in str(recurso.get("Aviso fecha")):
        fallo("shotlist %s: fechas %r" % (proy, {k: (v.get("Fecha grabación"), v.get("Aviso fecha")) for k, v in fila.items()}))


def ffmpeg_junto_a(ffprobe):
    """El ffmpeg de la misma carpeta que el ffprobe (en Windows, bin/ no está en el PATH); si no, el del PATH."""
    if ffprobe:
        p = Path(ffprobe)
        junto = p.with_name(p.name.lower().replace("ffprobe", "ffmpeg"))
        if junto.exists():
            return str(junto)
    return shutil.which("ffmpeg") or "ffmpeg"


def enlaces(proy, shotlist, mxf, clips, a):
    """ENLACES: los tres clips, el HDR marcado, el vertical derecho (centrado entre bandas) y el estéreo en dos pistas distintas."""
    if sorted(clips) != ["C0001", "GX010045", "OTRO DISCO"]:
        fallo("shotlist %s: clips %s, esperados los dos enlazados y el del otro disco" % (proy, clips))
    from openpyxl import load_workbook
    filas = list(load_workbook(shotlist, read_only=True).active.iter_rows(values_only=True))
    cab = filas[0]
    fila = {f[cab.index("Clip")]: dict(zip(cab, f)) for f in filas[1:]}
    if fila["C0001"].get("Resolución origen") != "1080x1920 (vertical)":
        fallo("shotlist %s: el vertical debe constar como 1080x1920 (vertical): %r" % (proy, fila["C0001"].get("Resolución origen")))
    if fila["GX010045"].get("Color origen") != "HLG":
        carpeta = sorted(c for c in glob.glob(str(Path(a.informe) / "Diagnostico_*")) if Path(c).is_dir())[-1]
        leido = [m for m in json.loads((Path(carpeta) / "mxf.json").read_text(encoding="utf-8")) if "GX010045" in m["ruta"]]
        print("Lo que leyó el análisis del GoPro:", json.dumps(leido, ensure_ascii=False))
        print("Fila del shotlist:", fila["GX010045"])
        fallo("shotlist %s: el GoPro HDR no sale como HLG (%r)" % (proy, fila["GX010045"].get("Color origen")))
    inicio = sum(int(x) * k for x, k in zip(fila["C0001"]["TC inicio (bruto)"].replace(";", ":").split(":"), (3600, 60, 1, 0.04)))
    r = subprocess.run([ffmpeg_junto_a(a.ffprobe), "-hide_banner", "-ss", "%.2f" % (inicio + 0.3), "-t", "1.2", "-i", str(mxf),
                        "-vf", "cropdetect=limit=24:round=2:reset=0", "-f", "null", "-"], capture_output=True, text=True)
    ancho = int(r.stderr.rsplit("crop=", 1)[1].split(":")[0]) if "crop=" in r.stderr else 0
    if not 590 <= ancho <= 630:
        fallo("bruto %s: el vídeo vertical no sale derecho y centrado (ancho de imagen %d, esperado ~608)" % (proy, ancho))
    vol = []
    for k in (0, 1):
        r = subprocess.run([ffmpeg_junto_a(a.ffprobe), "-hide_banner", "-ss", "%.2f" % (inicio + 0.3), "-t", "1", "-i", str(mxf),
                            "-map", "0:a:%d" % k, "-af", "astats=measure_overall=Peak_level:measure_perchannel=0", "-f", "null", "-"],
                           capture_output=True, text=True)
        vol.append(r.stderr)
    if "Peak level dB: -inf" not in vol[0] or "Peak level dB: -inf" in vol[1]:
        fallo("bruto %s: el estéreo del móvil (izquierdo mudo, derecho con tono) no sale como pista 1 muda y pista 2 con sonido" % proy)


def tc_shotlist(ruta):
    """(TC de inicio de los clips incluidos, proyecto principal, clip y fecha de cada fila en orden), del xlsx o, si
    no había openpyxl, del csv de respaldo. "15 SEPT (también en HONDURAS)" cuenta como 15 SEPT."""
    principal = lambda t: str(t or "").split(" (también en ")[0].strip()
    if ruta.suffix == ".xlsx":
        from openpyxl import load_workbook
        filas = list(load_workbook(ruta, read_only=True).active.iter_rows(values_only=True))
        col, colp, colc, colf, colg = (filas[0].index(x) for x in ("TC inicio (bruto)", "Proyecto", "Clip", "Fecha fichero", "Fecha grabación"))
        return ([f[col] for f in filas[1:] if f[col]], [principal(f[colp]) for f in filas[1:]], [f[colc] for f in filas[1:]],
                [str(f[colg] or f[colf] or "") for f in filas[1:]])
    import csv
    with open(ruta, encoding="utf-8-sig", newline="") as f:
        filas = list(csv.DictReader(f, delimiter=";"))
    return ([x["tc_in"] for x in filas if x.get("tc_in")], [principal(x.get("proyecto")) for x in filas], [x.get("clip") for x in filas],
            [x.get("fecha_grabacion") or x.get("fecha") or "" for x in filas])


if __name__ == "__main__":
    main()
