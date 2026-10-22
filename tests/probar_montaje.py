#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Prueba de la pantalla de montaje sin ventana, sobre el disco de prueba ya analizado.

  python tests/probar_montaje.py preparar CARPETA_ANALISIS CARPETA_BRUTOS
      Escribe montaje.json como si en la pantalla de montaje se hubiera:
        - cortado 15 SEPT CENTROAMERICA tras el primer clip,
        - invertido el orden de los dos clips del segundo bruto y llamado a ese bruto "PRUEBA MONTAJE",
        - quitado de ABORTO el clip 2018_21043491-c.
  python tests/probar_montaje.py comprobar CARPETA_BRUTOS
      Tras exportar con --usar-analisis --proyectos "15 SEPT CENTROAMERICA;ABORTO": exactamente esos brutos, con
      ese orden en el shotlist, el clip quitado fuera y la columna Montaje. Sale con 1 si no cuadra.
"""

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import exportar_brutos as eb  # noqa: E402
import montaje  # noqa: E402

SEPT, ABORTO = "15 SEPT CENTROAMERICA", "ABORTO"


def preparar(analisis, brutos):
    leer = lambda n: json.loads((Path(analisis) / n).read_text(encoding="utf-8"))
    clips, cruce, mxf = leer("clips.json"), leer("cruce.json"), leer("mxf.json")
    m = montaje.Montaje.desde_analisis(clips, cruce, mxf, json.loads(json.dumps(eb.CONFIG_DEFECTO)), brutos)
    clave = {c["nombre"]: c["clave"] for b in m.brutos(SEPT) + m.brutos(ABORTO) for c in b["clips"]}
    print("15 SEPT:", [[c["nombre"] for c in b["clips"]] for b in m.brutos(SEPT)])
    m.cortar(SEPT, clave["2017 07 21"])
    m.bajar(SEPT, [clave["2017 07 21"]])
    m.renombrar(SEPT, 1, "PRUEBA MONTAJE")
    m.quitar(ABORTO, [clave["2018_21043491-c"]])
    m.guardar(brutos)
    print("Montaje:", {p: [(b["nombre"], [c["nombre"] for c in b["clips"]]) for b in m.brutos(p)] for p in (SEPT, ABORTO)})


def shotlist(ruta):
    from openpyxl import load_workbook
    wb = load_workbook(ruta, read_only=True)
    filas = list(wb["Shotlist"].iter_rows(values_only=True))
    cab = filas[0]
    return [dict(zip(cab, f)) for f in filas[1:]], wb.sheetnames


def comprobar(brutos):
    fallos = []
    res = json.loads((Path(brutos) / "brutos_resultado.json").read_text(encoding="utf-8"))
    hechos = {b["proyecto"]: b for b in res if b.get("proyecto_avid") in (SEPT, ABORTO)}
    esperado = {SEPT + " SEPTIEMBRE DE 2021": ["2017 07 13"],
                "PRUEBA MONTAJE": ["2017 07 21 ESPAÑA", "2017 07 21"],
                SEPT + " PIEZAS SEPTIEMBRE DE 2021": ["PIEZA 15 SEPT"],
                ABORTO + " SEPTIEMBRE DE 2021": ["2018-10-30T013549Z"]}
    if sorted(hechos) != sorted(esperado):
        fallos.append("brutos %s, esperados %s" % (sorted(hechos), sorted(esperado)))
    for nombre, clips in esperado.items():
        b = hechos.get(nombre)
        if not b:
            continue
        if not str(b.get("estado", "")).startswith("OK"):
            fallos.append("%s: estado %s" % (nombre, b.get("estado")))
        filas, hojas = shotlist(Path(brutos) / b["shotlist"])
        if [f["Clip"] for f in filas] != clips:
            fallos.append("%s: orden %s, esperado %s" % (nombre, [f["Clip"] for f in filas], clips))
        if nombre == "PRUEBA MONTAJE" and [f.get("Montaje") for f in filas] != ["automático", "movido a mano"]:
            fallos.append("%s: columna Montaje %s" % (nombre, [f.get("Montaje") for f in filas]))
        if nombre.startswith(ABORTO) and "Quitados" not in hojas:
            fallos.append("%s: falta la hoja Quitados" % nombre)
    if fallos:
        print("FALLO montaje:\n  " + "\n  ".join(fallos))
        sys.exit(1)
    print("Montaje OK:", {n: hechos[n]["incluidos"] for n in esperado})


if __name__ == "__main__":
    if sys.argv[1] == "preparar":
        preparar(sys.argv[2], sys.argv[3])
    else:
        comprobar(sys.argv[2])
